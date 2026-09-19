"""300-offtaker synthetic rehearsal; no Norwich customer data or external I/O."""
import io
import time
import pytest
from openpyxl import Workbook
from sqlalchemy import select
from api.db import SessionLocal
from api.models import BillingReportSubscription
from tests.test_offtaker_upload import _make_tenant, _make_array_with_bill, _bulk_import

B="/v1/array-operator/billing"


def test_300_row_messy_roster_preview_commit_and_replay(client,monkeypatch,tmp_path):
    monkeypatch.delenv("ANTHROPIC_API_KEY",raising=False)
    tid,auth=_make_tenant()
    arrays=[_make_array_with_bill(tid,f"Audit Array {i:03d}",f"GMP-SCALE-{i:03d}",with_bill=True) for i in range(100)]
    wb=Workbook(); ws=wb.active
    ws.append(["Norwich synthetic readiness rehearsal: 300 accounts"])
    ws.append([])
    ws.append(["Customer","Array","% Allocation","Rate ($/kWh)","Contact e-mail"])
    for i in range(300):
        ws.append([f"Offtaker {i:03d}",f"Audit Array {i//3:03d}",[25,0.25,"25%"][i%3],"$0.18398",f"OFFTAKER{i:03d}@Example.COM"])
    ws.append(["Total",None,7500])
    data=io.BytesIO(); wb.save(data)
    start=time.monotonic()
    response=_bulk_import(client,auth,"norwich-300.xlsx",data.getvalue(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert response.status_code==200,response.text
    preview=response.json()
    assert len(preview["rows"])==300,preview["summary"]
    with SessionLocal() as db:
        assert not db.execute(select(BillingReportSubscription).where(BillingReportSubscription.tenant_id==tid)).scalars().all()
    rows=[]
    for i,r in enumerate(preview["rows"]):
        assert r["offtaker_name"]==f"Offtaker {i:03d}"
        assert r["allocation_pct"]==pytest.approx(0.25),r
        assert r["net_rate_per_kwh"]==pytest.approx(0.18398),r
        assert r["discount_pct"] is None,r
        assert r["email"]==f"offtaker{i:03d}@example.com"
        assert r["matched_array_id"]==arrays[i//3][0],r
        assert r["matched_utility_account_id"]==arrays[i//3][1],r
        rows.append({"offtaker_name":r["offtaker_name"],"array_id":r["matched_array_id"],"utility_account_id":r["matched_utility_account_id"],"allocation_pct":r["allocation_pct"],"net_rate_per_kwh":r["net_rate_per_kwh"],"email":r["email"]})
    body={"rows":rows+[dict(rows[0])],"cadence":"monthly","delivery_mode":"approval"}
    result=client.post(B+"/subscriptions/bulk-commit",json=body,headers={"Authorization":auth})
    assert result.status_code==200,result.text
    created=result.json()
    assert created["created"]==300,created
    assert not created["failed"],created
    assert len(created["skipped"])==1
    replay=client.post(B+"/subscriptions/bulk-commit",json=body,headers={"Authorization":auth}).json()
    assert replay["created"]==0 and len(replay["skipped"])==301 and not replay["failed"],replay
    with SessionLocal() as db:
        subs=db.execute(select(BillingReportSubscription).where(BillingReportSubscription.tenant_id==tid)).scalars().all()
        assert len(subs)==300
        assert all(s.net_rate_per_kwh==pytest.approx(0.18398) and s.allocation_pct==0.25 and s.delivery_mode=="approval" and s.last_sent_at is None for s in subs)
    print(f"300-row preview + commit + replay seconds: {time.monotonic()-start:.3f}")
    from datetime import date
    from api.billing import delivery, invoice, payments
    from openpyxl import load_workbook
    import fitz
    # Independent arithmetic for every account: 5000 kWh * 25% * $0.18398 * 90%.
    # Verify the shared render payload, email amount, and Stripe cents agree.
    with SessionLocal() as db:
        subs=db.execute(select(BillingReportSubscription).where(BillingReportSubscription.tenant_id==tid).order_by(BillingReportSubscription.id)).scalars().all()
        for index,sub in enumerate(subs):
            match=delivery.build_match(sub)
            assert match.computed_invoice["amount_owed"]==206.98,match.computed_invoice
            assert match.computed_invoice["rate_is_operator_entered"] is True
            payload=invoice.invoice_for_period(match,match.latest_period,date.today())
            assert payload["amount_owed"]==206.98
            assert payments._amount_cents_from_match(match)==20698
            _,html,text=delivery._email_html(match,sub,is_test=False)
            assert "$206.98" in html and "$206.98" in text
            if index==0:
                pdf=tmp_path/"invoice.pdf";xlsx=tmp_path/"invoice.xlsx"
                invoice.render_invoice_pdf(match,pdf)
                invoice.render_invoice_xlsx(match,xlsx)
                with fitz.open(pdf) as doc:
                    assert "$206.98" in "".join(page.get_text() for page in doc)
                book=load_workbook(xlsx,data_only=True)
                assert any(cell.value in (206.98,"$206.98") for sheet in book for row in sheet for cell in row)
    print(f"300-row arithmetic + email/Stripe checks and representative PDF/XLSX seconds: {time.monotonic()-start:.3f}")


def test_foreign_subscription_and_draft_routes_reject_access(client):
    from tests.test_billing_delivery import _make_tenant as make, _upload
    tid,auth=make(); other,foreign=make()
    sid=_upload(client,auth,"norwich.xlsx").json()["subscription"]["id"]
    own={"Authorization":auth}; bad={"Authorization":foreign}
    draft=client.post(f"{B}/subscriptions/{sid}/draft",headers=own).json()
    did=(draft.get("draft") or draft)["id"]
    requests=[("get",f"/subscriptions/{sid}/preview",None),("get",f"/subscriptions/{sid}/payments",None),("post",f"/subscriptions/{sid}/send-now",None),("patch",f"/subscriptions/{sid}",{"client_email":"intruder@example.test"}),("post",f"/drafts/{did}/approve",None),("post",f"/drafts/{did}/test",None)]
    for method,path,body in requests:
        kwargs={"headers":bad}
        if body is not None: kwargs["json"]=body
        r=getattr(client,method)(B+path,**kwargs)
        assert r.status_code==404,(method,path,r.status_code,r.text)
